"""Finish training-anatomy selection, frozen paired evaluation, paper notes, and media."""
import hashlib
import json
import os
from pathlib import Path
import signal
import subprocess
import sys
import time

ROOT=Path(__file__).resolve().parents[1]
V=ROOT/'research/validation/EXP0094_PPO_NAV_20261011'
R=ROOT/'research/runs/EXP0094_PPO_NAV_20261011/seed0'
MEDIA=Path('/home/wj/桌面/VRL_RL_局部导航_训练测试_GIF_20261011')

def read_rows(file):
    rows=[]
    if not file.exists():return rows
    for l in file.open():
        try:rows.append(json.loads(l))
        except json.JSONDecodeError:pass
    return rows

def run(args,log):
    with log.open('a') as f:
        subprocess.run([sys.executable,'-u','-m',*args],cwd=ROOT,stdout=f,stderr=subprocess.STDOUT,check=True)

def document(checkpoint,summary):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();d.styles['Normal'].font.name='Noto Sans CJK SC';d.styles['Normal'].font.size=Pt(10.5)
    d.add_heading('VRL：强化学习局部导航方法修订与实际实验记录',0)
    d.add_paragraph('V3 的主方法为监督运动预测与固定代价候选选择，没有 RL。本次保留成功率最高的测地分配 + 固定 TPG 全局组合，将局部方向和速度决策替换为 PPO 策略。旧文档的成功率不能直接作为 RL 新方法的成绩。')
    d.add_heading('实际使用 RL 的位置',1)
    d.add_paragraph('延迟图像及已发指令历史、术前地图和冻结运动预测器的候选动作后果输入循环 actor。actor 输出候选指令的概率，确定运动方向与速度；在网络后不再运行固定代价 argmin 决定动作。原规则与代价作为初始化和建议特征，策略打分参数与循环头由环境奖励的 PPO 优化。壁护、TPG 等待和间距护盾保留。')
    d.add_paragraph('适配通过测量历史和预测信念进行，覆盖未知血流、响应漂移和成像延迟。仿真真实位置及剩余质量不进入策略输入；它们只计算训练奖励和评估指标。这里没有把监督预测器改名为 RL。')
    d.add_heading('训练和检查点选择',1)
    cfg=json.load(open(R/'config.json'));complete=json.load(open(R/'TRAINING_COMPLETE.json'))
    d.add_paragraph(f"已完成 {complete['iterations']} 次 PPO 更新、{complete['agent_steps']:,} 次有效智能体决策。训练只使用九种训练解剖，N=1–3、三档入口流速、1–3 步延迟和随机动力学扰动；五种留出解剖未用于训练或选择。预测器冻结。所选检查点：{checkpoint}。检查点仅按独立训练解剖调参池选择。")
    d.add_paragraph('早期未排除不稳定 CPU 6/7 的尝试发生底层崩溃，已单独归档并作废。本文记录只使用重新初始化并按仓库既有设置绑核的正式实验。')
    d.add_heading('252 个完整配对困难场景的结果',1)
    table=d.add_table(rows=1,cols=7);table.style='Light Shading Accent 1'
    for cell,text in zip(table.rows[0].cells,['方法','场景数','Strict %','清除 %','T90_300 s','壁 robot·s','RL 输出改变 %']):cell.text=text
    for name,s in summary['results'].items():
        row=table.add_row().cells
        for c,text in zip(row,[name,str(s['n']),f"{s['strict_success']*100:.1f}",f"{s['removal']*100:.1f}",f"{s['T90_300']:.1f}",f"{s['wall_contact_s']:.2f}",f"{s['rl_output_changed_fraction']*100:.3f}"]):c.text=text
    name='seed0/'+Path(checkpoint).name;pair=summary['paired_to_parent'][name]
    ci=pair['CI95_pp'];d.add_paragraph(f"相对原非 RL 方法的配对 Strict 变化：{pair['strict_delta_pp']:+.2f} 个百分点，95% CI [{ci[0]:+.2f}, {ci[1]:+.2f}]。scenario_hash 不一致数为 {pair['scenario_hash_mismatches']}。")
    if ci[0]<=0:
        d.add_paragraph('该结果尚不能支持 RL 显著优于原方法。应如实说明：RL 导航已经实现并接受真实奖励训练，但其性能提升需要更多证据；不得把使用 RL 和 RL 带来提升混为一谈。')
    d.add_paragraph('初始化策略另列为对照，用于区分策略初始化/安全限制的影响和 PPO 更新的影响。报告中的 rl_update_command_changed_fraction 单独记录相对未训练打分器的实际方向/速度变化。失败回合完整计入，不能只展示成功动画。当前是一个 RL 种子的初版，不声称三种子稳定性。')
    measured=summary['results'][name]
    d.add_paragraph(f"选定策略共执行 {measured['rl_decisions']:,} 次局部 RL 决策，相对未训练打分器的实际方向/速度变化占 {100*measured['rl_update_command_changed_fraction']:.4f}%；经过壁护后，相对原非 RL 方法的局部输出变化占 {100*measured['rl_output_changed_fraction']:.4f}%。本版策略动作偏离原方法的比例很小，不能据此声称已证明 RL 的适配收益。")
    d.add_paragraph('以上场景属于原 V3 开发评估矩阵，尚未执行封存测试。若配对自举区间为 [0, 0]，它表示该有限场景集中 Strict 没有不一致，不构成对未来场景或更多训练种子的等价性证明。分条件和解剖组统计另存于 full_report/condition_topology_breakdown.csv；动作归因见 full_report/ATTRIBUTION_AND_SCENARIOS.md。')
    d.add_heading('可复核产物',1)
    for text in [str(V/'full_hard/episodes.jsonl'),str(V/'full_report/REPORT.md'),str(R/'training.jsonl'),str(R/'training_traces'),str(MEDIA/'index.html')]:d.add_paragraph(text)
    target=ROOT/'VRL_RL_局部导航_方法修订与实验记录_20261011.docx';d.save(target)
    return target

def resume_legacy_media():
    p=Path('/home/wj/桌面/VRL_V3_全部场景_3D_GIF_20261011');data=json.load(open(p/'batch_pid.json'));pid=data['pid']
    try:
        if os.getpgid(pid)!=pid:return
        affinity=set(range(24))-{6,7}
        for proc in Path('/proc').iterdir():
            if proc.name.isdigit():
                try:
                    if os.getpgid(int(proc.name))==pid:os.sched_setaffinity(int(proc.name),affinity)
                except (ProcessLookupError,PermissionError):pass
        os.killpg(pid,signal.SIGCONT)
        progress=json.load(open(p/'progress.json'));progress.update(paused=False,role='旧 V3 非 RL 方法、外部方法和消融的对照媒体')
        (p/'progress.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2))
    except ProcessLookupError:pass

def main():
    os.chdir(ROOT);V.mkdir(parents=True,exist_ok=True)
    os.sched_setaffinity(0,os.sched_getaffinity(0)-{6,7})
    for name in ('OMP_NUM_THREADS','MKL_NUM_THREADS','OPENBLAS_NUM_THREADS'):os.environ[name]='1'
    while len(read_rows(V/'tune_final/episodes.jsonl'))<324:
        print('Waiting for the 324 training-anatomy tuning records',flush=True);time.sleep(30)
    from scripts.report_predictive_nav_rl import report
    tuned=report([V/'tune_final/episodes.jsonl'],V/'tune_final_report')
    candidate={name:s for name,s in tuned['results'].items() if name.startswith('seed0/policy_it') and s['n']==108}
    if not candidate:raise RuntimeError('No complete trained checkpoint in tuning results')
    best=max(candidate,key=lambda name:(candidate[name]['strict_success'],-candidate[name]['wall_contact_s'],-candidate[name]['T90_300'],name))
    checkpoint=str((R/Path(best).name).relative_to(ROOT));initial=str((R/'initial.pt').relative_to(ROOT))
    selection=dict(checkpoint=checkpoint,selection_pool='108 training-anatomy tuning cases',tuning_scores=candidate,
                   sha256=hashlib.sha256((ROOT/checkpoint).read_bytes()).hexdigest(),selected_at=time.strftime('%F %T'))
    (V/'SELECTED.json').write_text(json.dumps(selection,ensure_ascii=False,indent=2));print('Selected',checkpoint,flush=True)
    run(['scripts.eval_exp0090','--out',str(V/'full_hard'),'--protocol','hard','--methods',
         f'rl_nav={checkpoint},rl_nav={initial},rl_nav_off={initial}','--workers','22'],V/'full_hard.log')
    if len(read_rows(V/'full_hard/episodes.jsonl'))!=756:raise RuntimeError('Full paired evaluation is incomplete; inspect failures')
    summary=report([V/'full_hard/episodes.jsonl'],V/'full_report')
    if any(p['scenario_hash_mismatches'] for p in summary['paired_to_parent'].values()):raise RuntimeError('Paired geometry hashes disagree')
    doc=document(checkpoint,summary)
    from scripts.export_predictive_rl_media import make_tasks
    tasks=make_tasks(checkpoint,MEDIA,True)
    indexed={(r['method']+'='+r['method_arg'],r['cell'],r['anatomy'],r['clusters'],r['seed']):r for r in read_rows(V/'full_hard/episodes.jsonl')}
    for t in tasks:
        an,n,s,*_=t['job'];t['reference']=indexed.get((t['method'],t['cell'],an,n,s))
    from scripts.export_vrl_v3_all_gifs import atomic_json
    atomic_json(MEDIA/'tasks.json',tasks)
    (MEDIA/'说明.md').write_text(f'当前主方法为 PPO 强化学习局部导航，检查点 {checkpoint}。真实 PPO 训练轨迹与评估回放分别保存。每份 GIF 同时保存完整三维轨迹、参数、指标和来源。\n\n训练只使用九种训练解剖；困难矩阵包含全部 252 个场景，分别导出训练后、未训练初始化和关闭 RL 对照；弱流 V5 导出全部 1260 个场景的训练后及关闭 RL 回放。成功和失败都保留。旧 V3 非 RL 和外部方法动画位于旁边的 VRL_V3_全部场景_3D_GIF_20261011 文件夹，作为对照材料，不能称作 RL 主方法。\n\n完整困难矩阵报告：{V/"full_report/REPORT.md"}\n方法修订与实际实验记录：{doc}\n动画使用 PyBullet 三维血管网格与实际轨迹；帧数最多 180，时间加速但覆盖完整回合。\n')
    atomic_json(V/'ALGORITHM_COMPLETE.json',dict(checkpoint=checkpoint,model_doc=str(doc),full_paired_cases=252,media_expected=len(tasks)))
    print('Trained RL algorithm and paired evaluation complete; exporting all media',flush=True)
    # The original GIF scope is retained as comparison material. Pin the
    # suspended workers before restoring its finite, resumable export.
    resume_legacy_media()
    run(['scripts.export_predictive_rl_media','--out',str(MEDIA),'--workers','18'],V/'media.log')
    atomic_json(V/'MEDIA_BATCH_FINISHED.json',json.load(open(MEDIA/'progress.json')))

if __name__=='__main__':main()
