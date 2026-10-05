"""Convert the Markdown draft to a Word document (headings, paragraphs, bullets, tables, inline bold/code).

usage: md_to_docx.py IN.md OUT.docx [--title "..."]
Handles the subset of Markdown used by research/MANUSCRIPT_DRAFT_*.md. A CJK font is set on every run so
Word does not fall back to a Latin font for Chinese text.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CJK = '微软雅黑'
MONO = 'Consolas'


def style_run(run, font=CJK):
    run.font.name = font
    run.font.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)


def add_rich(par, text, font=CJK):
    """Inline **bold** and `code`."""
    for piece in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            r = par.add_run(piece[2:-2]); r.bold = True; style_run(r, font)
        elif piece.startswith('`') and piece.endswith('`'):
            r = par.add_run(piece[1:-1]); style_run(r, MONO); r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
        else:
            r = par.add_run(piece); style_run(r, font)


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def add_table(doc, rows):
    header, align, body = rows[0], rows[1], rows[2:]
    t = doc.add_table(rows=len(body)+1, cols=len(header))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    right = [i for i, a in enumerate(align) if a.endswith(':')]
    for j, h in enumerate(header):
        cell = t.cell(0, j); cell.text = ''
        add_rich(cell.paragraphs[0], h)
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for i, row in enumerate(body, 1):
        for j, v in enumerate(row[:len(header)]):
            cell = t.cell(i, j); cell.text = ''
            p = cell.paragraphs[0]; add_rich(p, v)
            if j in right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return t


def convert(md: str, out: Path, title=None, src_dir=Path('.')):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = CJK; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    if title:
        doc.add_heading(title, 0)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        if line.startswith('---') and set(line.strip()) == {'-'}:
            doc.add_paragraph(); i += 1; continue
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', line.strip())
        if m:                                           # figure: image + italic caption
            path = Path(m.group(2))
            if not path.is_absolute():
                path = src_dir/path
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(cap, m.group(1))
            for r in cap.runs:
                r.italic = True; r.font.size = Pt(9)
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            h = doc.add_heading('', min(len(m.group(1)), 4))
            add_rich(h, m.group(2)); i += 1; continue
        if line.lstrip().startswith('|') and i+1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1]):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows); doc.add_paragraph(); continue
        m = re.match(r'^(\s*)([-*])\s+(.*)$', line)
        if m:
            depth = len(m.group(1))//2
            p = doc.add_paragraph(style='List Bullet' if depth == 0 else 'List Bullet 2')
            add_rich(p, m.group(3)); i += 1; continue
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if m:
            depth = len(m.group(1))//2
            p = doc.add_paragraph(style='List Number' if depth == 0 else 'List Number 2')
            add_rich(p, m.group(3)); i += 1; continue
        if line.lstrip().startswith('>'):
            p = doc.add_paragraph(style='Intense Quote'); add_rich(p, line.lstrip('> ')); i += 1; continue
        # paragraph: join continuation lines
        buf = [line]
        while i+1 < len(lines) and lines[i+1].strip() and not re.match(r'^\s*(#|\||[-*]\s|\d+\.\s|>)', lines[i+1]) \
                and not (lines[i+1].startswith('---') and set(lines[i+1].strip()) == {'-'}):
            i += 1; buf.append(lines[i].strip())
        p = doc.add_paragraph(); add_rich(p, ''.join(buf) if any('一' <= ch <= '鿿' for ch in buf[0]) else ' '.join(buf))
        i += 1
    doc.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('src', type=Path); ap.add_argument('dst', type=Path); ap.add_argument('--title')
    a = ap.parse_args()
    convert(a.src.read_text(), a.dst, a.title, a.src.parent)
    print('wrote', a.dst)


if __name__ == '__main__':
    main()
